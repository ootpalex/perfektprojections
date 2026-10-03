"""metadata_inputs.py - build 25 Metadata's input tabs from StatsPlus, no paste.

Writes the nine input CSVs that engine/metadata_calibrate.py turns into the
league's metadata Data Points (the constants Sync Metadata pushes into The Sheet
Hitters / Pitchers). Sources, per tab:

  public stats API (playerbatstatsv2 / playerpitchstatsv2 / playerfieldstatsv2,
  splits 1 overall, 2 vs L, 3 vs R; rows summed per player across team stints):
      Hitting_Data.csv    every player with a PA in the season
      Pitching_Data.csv   every pitcher with a BF in the season
      Fielding_Data.csv   one table per position C..RF, one row per player-position
  API + the ratings of the last live pull of the regular season (on or before
  in-game <season>-10-01, from the ratings archive; the newest pull
  .cache/statsplus_<slug>.json when the archive has none):
      Batter_Ratings.csv  vR table = hitters with PA vs RHP, vL table = PA vs LHP
      Fielding_Ratings.csv  position players: ratings, height, total fielding IP
      SP_Ratings.csv / RP_Ratings.csv  the pitchers in your SP / RP paste, BF vs R / vs L
  your paste in 25 Metadata.xlsx (StatsPlus has no stats split by role):
      SP_Data.csv / RP_Data.csv  <- tabs 'SP Data' / 'RP Data'

Stages (the Recalibrate bat runs them in order with a paste reminder between):
  --stage auto    the five API/pull tabs; ends with the reminder
  --stage roles   the two paste tabs + the two pitcher ratings tabs, after
                  reconciling your paste against the API season totals. A
                  paste of another season stops, even with --accept-paste; the
                  message names the season the paste is from when one of the
                  two seasons before matches it (paste_season)
  --stage all     both (default)

Usage:
  python tgs-viz/ingest/metadata_inputs.py --league BLM --out tgs-viz/engine/calib/BLM/metadata_inputs
  python tgs-viz/engine/metadata_calibrate.py --inputs-dir <out> --json <calib>/metadata-latest.json
  python tgs-viz/ingest/sync_datapoints.py --league BLM --metadata-calib <json> --calib <constants> --write

Season: the in-game date decides (Oct-Dec = this year's season, Jan-Mar = last
year's); April-September refuses unless --allow-partial. --year overrides.
Exactness: every stat is the API value; IP is written in OOTP thirds notation
so the port's DOLLARDE path matches the workbook; heights are the pull's cm
rendered as feet-inches (what OOTP shows). Nothing is estimated.

StatsPlus: requests carry the league's saved token on their own (statsplus.py).
/teams is reused from the date-keyed cache while the in-game date is unchanged,
so the two bat stages download it once. The stat feeds are saved per season in
.cache/metadata_<slug>_<season>/; a feed is saved only when every column this
script sums is there and it has rows, and a saved feed that fails that check
is downloaded again.

Engine boundary: a league with leagues.<LG>.engine_first_season in settings refuses
a season before it (older seasons were played on another OOTP engine), before the
/teams and stat-feed requests; with --year, before any request.

Exit codes: 1 = a stop with a message (season, paste or workbook problem);
2 = bad arguments; 3 = StatsPlus refused a request or sent something that is
not the data. On 3 no input CSV or manifest is written, and no bad feed is
saved.
"""
import argparse
import csv
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
import statsplus as S  # noqa: E402
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "tools"))
import settings as ST  # noqa: E402

LEAGUE_DIRS = {"TGS": "The Sheets TGS", "BLM": "The Sheets BLM"}
SLUGS = {lg: s for lg, s in ST.slug_map().items() if lg in LEAGUE_DIRS}
PITCHER_POS = {"SP", "RP", "CL", "P"}
POS_CODE = {"P": 1, "C": 2, "1B": 3, "2B": 4, "3B": 5, "SS": 6, "LF": 7, "CF": 8, "RF": 9, "DH": 10,
            "SP": 1, "RP": 1, "CL": 1}
FIELD_POS = {2: "C", 3: "1B", 4: "2B", 5: "3B", 6: "SS", 7: "LF", 8: "CF", 9: "RF"}
PASTE_TABS = {"SP Data": "SP_Data.csv", "RP Data": "RP_Data.csv"}
DERIVED_TABS = {"Batter Ratings": "Batter_Ratings.csv", "Fielding Ratings": "Fielding_Ratings.csv",
                "SP Ratings": "SP_Ratings.csv", "RP Ratings": "RP_Ratings.csv"}
STAT_TABS = {"Hitting Data": "Hitting_Data.csv", "Pitching Data": "Pitching_Data.csv",
             "Fielding Data": "Fielding_Data.csv"}

HIT_HDR = ["ID", "Name", "ORG", "G", "GS", "PA", "AB", "H", "1B", "2B", "3B", "HR", "RBI", "R", "BB",
           "IBB", "HP", "SH", "SF", "CI", "SO", "GIDP", "PI/PA", "SB", "CS", "UBR"]
PIT_HDR = ["ID", "Name", "ORG", "G", "GS", "W", "L", "IP", "BF", "AB", "HA", "1B", "2B", "3B", "HR", "R",
           "ER", "BB", "IBB", "K", "HP", "SH", "SF", "WP", "BK", "CI", "DP", "IR", "IRS", "PI", "GB", "FB",
           "SB", "CS"]
FLD_HDR = ["ID", "Name", "ORG", "POS", "G", "GS", "TC", "A", "PO", "E", "DP", "TP", "ZR", "SBA", "RTO",
           "IP", "PB", "CER", "BIZ-R", "BIZ-Rm", "BIZ-L", "BIZ-Lm", "BIZ-E", "BIZ-Em", "BIZ-U", "BIZ-Um",
           "BIZ-Z", "BIZ-Zm", "BIZ-I", "FRM", "ARM"]
BAT_RAT_HDR = ["ID", "POS", "Name", "PA", "B", "BA vL", "GAP vL", "POW vL", "EYE vL", "K vL", "BA vR",
               "GAP vR", "POW vR", "EYE vR", "K vR", "SPE", "STE", "RUN", "SR"]
PIT_RAT_HDR = ["ID", "POS", "Name", "BF", "T", "STU vL", "HRR vL", "PBABIP vL", "CON vL", "STU vR",
               "HRR vR", "PBABIP vR", "CON vR", "HLD"]
FLD_RAT_HDR = ["ID", "POS", "Name", "IP", "HT", "C ABI", "C FRM", "C ARM", "IF RNG", "IF ERR", "IF ARM",
               "TDP", "OF RNG", "OF ERR", "OF ARM"]

BAT_SUM = ["pa", "ab", "h", "d", "t", "hr", "rbi", "r", "bb", "ibb", "hp", "sh", "sf", "ci", "k", "gdp",
           "sb", "cs", "ubr", "g", "gs", "pitches_seen"]
PIT_SUM = ["g", "gs", "w", "l", "outs", "bf", "ab", "ha", "da", "ta", "hra", "r", "er", "bb", "iw", "k",
           "hp", "sh", "sf", "wp", "bk", "ci", "dp", "ir", "irs", "pi", "gb", "fb", "sb", "cs"]
FLD_SUM = ["g", "gs", "tc", "a", "po", "e", "dp", "tp", "zr", "sba", "rto", "ip", "ipf", "pb", "er",
           "framing", "arm"] + [f"opps_{i}" for i in range(6)] + [f"opps_made_{i}" for i in range(6)]


# ---------------------------------------------------------------- helpers

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def cell(v):
    """CSV cell text: integral floats as ints, other floats as repr, None as ''."""
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def thirds(outs):
    """Outs -> OOTP 'x.y' innings notation (y = thirds), as the workbook pastes."""
    outs = int(round(outs))
    return f"{outs // 3}.{outs % 3}"


def feet_inches(cm):
    """Pull height (cm) -> OOTP's display string 6' 4\" (rounded to the inch)."""
    cm = num(cm)
    if cm <= 0:
        return ""
    inches = int(round(cm / 2.54))
    return f"{inches // 12}' {inches % 12}\""


def season_of(date_str, year_arg, allow_partial):
    y, m = int(date_str[:4]), int(date_str[5:7])
    if year_arg:
        return year_arg
    if m <= 3:
        return y - 1
    if m <= 9 and not allow_partial:
        sys.exit(f"in-game date {date_str}: the regular season is still running. "
                 f"Run this after the season ends (October), or pass --year/--allow-partial.")
    return y


def engine_boundary_problem(league, season, first, version=None):
    """Why `season` can not feed `league`'s metadata, or None. first = the league's
    engine_first_season (settings), None when it has no boundary. Seasons before it
    were played on another OOTP engine: their rates are not this engine's."""
    if first is None or season >= first:
        return None
    on = f" (OOTP {version})" if version else ""
    return (f"STOP: {league} season {season} was played before this league's engine boundary: "
            f"leagues.{league}.engine_first_season = {first}{on}. Metadata from an older engine "
            f"would put the wrong rates into the calibration. Build season {first} or later "
            f"(--year), or correct engine_first_season in settings.local.json if the boundary is wrong.")


def check_engine_boundary(league, season):
    """Stop (exit 1, with the message) when `season` is before the league's boundary."""
    lg = ST.league(league) or {}
    msg = engine_boundary_problem(league, season, ST.engine_first_season(league), lg.get("ootp_version"))
    if msg:
        sys.exit(msg)


def aggregate(rows, key, cols):
    """Sum numeric columns per key across team stints; keep the first row's ids."""
    out = {}
    for r in rows:
        k = key(r)
        d = out.get(k)
        if d is None:
            d = out[k] = {c: 0.0 for c in cols}
            d["_team"] = r.get("team_id")
            d["_pos"] = r.get("position")
        for c in cols:
            d[c] += num(r.get(c))
    return out


def write_csv(path, header, rows, banner="Player List"):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([banner] + [""] * (len(header) - 1))
        w.writerow(header)
        for r in rows:
            w.writerow([cell(v) for v in r])


def write_side_by_side(path, header, left, right, tags=("vR", "vL")):
    """Two tables side by side with the workbook's '<--', '', '-->' spacer."""
    spacer = ["<--", "", "-->"]
    n = len(header)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Player List"] + [""] * (n - 1) + [tags[0], "", tags[1]] + ["Player List"] + [""] * (n - 1))
        w.writerow(header + spacer + header)
        for i in range(max(len(left), len(right))):
            a = [cell(v) for v in left[i]] if i < len(left) else [""] * n
            b = [cell(v) for v in right[i]] if i < len(right) else [""] * n
            w.writerow(a + ["", "", ""] + b)


def write_fielding(path, header, tables):
    """Eight position tables side by side, one blank column between (the port
    stops a table at the first blank header)."""
    n = len(header)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        top, hdr = [], []
        for i in range(len(tables)):
            top += ["Player List"] + [""] * (n - 1) + ([""] if i < len(tables) - 1 else [])
            hdr += header + ([""] if i < len(tables) - 1 else [])
        w.writerow(top)
        w.writerow(hdr)
        depth = max(len(t) for t in tables)
        for i in range(depth):
            row = []
            for ti, t in enumerate(tables):
                row += ([cell(v) for v in t[i]] if i < len(t) else [""] * n)
                if ti < len(tables) - 1:
                    row.append("")
            w.writerow(row)


def read_workbook_tab(xlsx, tab):
    """A pasted input tab: header row 2, rows from 3 while column A is filled.
    Returns (headers, rows) with Excel's cached values (openpyxl read-only)."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if tab not in wb.sheetnames:
        wb.close()
        sys.exit(f"{os.path.basename(xlsx)} has no '{tab}' tab")
    ws = wb[tab]
    it = ws.iter_rows(min_row=2, values_only=True)
    header = list(next(it))
    while header and header[-1] is None:
        header.pop()
    rows = []
    for r in it:
        if not r or not _is_player_id(r[0]):
            continue          # blank rows and OOTP's export footer line
        rows.append(list(r[:len(header)]))
    wb.close()
    return header, rows


def _is_player_id(v):
    """True for a real player id in column A. OOTP's list export ends with a
    footer line ('Wednesday, ... - OOTP Baseball 27.4 Build 76') that lands in
    column A when the list is pasted; it is not a player."""
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return v > 0
    return str(v).strip().isdigit() and int(str(v).strip()) > 0


def write_workbook_tab(xlsx, tab, path):
    """Copy a pasted tab to CSV (values only; layout as pasted). Rows 1-2 are the
    banner and headers; below them only rows that start with a player id are kept
    (single-table tabs). Side-by-side ratings tabs are copied whole."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[tab]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    wb.close()
    single = tab in PASTE_TABS
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for i, r in enumerate(rows):
            r = list(r)
            while r and r[-1] is None:
                r.pop()
            if not r:
                continue
            if single and i >= 2 and not _is_player_id(r[0]):
                continue
            w.writerow([cell(v) for v in r])
    return rows


# ---------------------------------------------------------------- fetch + cache

EXIT_REFUSED = 3      # StatsPlus refused, or its reply was not the data
# The columns build_stats reads from each feed. A feed without one is never
# saved: the tabs would get zeros in place of the season's numbers.
FEED_COLS = {"bat": ["player_id"] + BAT_SUM, "pit": ["player_id"] + PIT_SUM,
             "fld": ["player_id", "position"] + FLD_SUM}
FEED_ENDPOINT = {"bat": "playerbatstatsv2", "pit": "playerpitchstatsv2", "fld": "playerfieldstatsv2"}


def feed_problem(name, rows):
    """Why a stat feed can not be used, or None: no rows, or a column that
    build_stats reads is missing."""
    if not isinstance(rows, list) or not rows or not isinstance(rows[0], dict):
        return "no rows"
    missing = [c for c in FEED_COLS[name[:3]] if c not in rows[0]]
    if missing:
        return (f"no {', '.join(missing[:6])}{' ...' if len(missing) > 6 else ''} "
                f"column{'s' if len(missing) > 1 else ''}")
    return None


def _save_json(path, obj):
    """json.dump through a temp file, so a failed write keeps the old copy."""
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f)
    os.replace(tmp, path)


def fetch_season(base, season, cache_dir, refresh):
    """All stat feeds for the season, saved as the parsed rows per feed.
    A feed is saved only when feed_problem() finds nothing wrong, and a saved
    feed that fails the check is downloaded again. Raises StatsPlusRefused
    when StatsPlus refuses or sends a feed that can not be used."""
    os.makedirs(cache_dir, exist_ok=True)
    feeds = {"bat1": (S.fetch_batting, 1), "bat2": (S.fetch_batting, 2), "bat3": (S.fetch_batting, 3),
             "pit1": (S.fetch_pitching, 1), "pit2": (S.fetch_pitching, 2), "pit3": (S.fetch_pitching, 3),
             "fld1": (S.fetch_fielding, 1)}
    out = {}
    for name, (fn, split) in feeds.items():
        p = os.path.join(cache_dir, f"{name}.json")
        saved = None
        if os.path.exists(p) and not refresh:
            try:
                saved = json.load(open(p, encoding="utf-8"))
            except ValueError:
                saved = None
            why = feed_problem(name, saved)
            if why:
                print(f"  {name}: the saved copy is not usable ({why}); downloading it again")
                saved = None
        if saved is not None:
            out[name] = saved
        else:
            rows = fn(base, year=season, split=split)
            why = feed_problem(name, rows)
            if why:
                raise S.StatsPlusRefused("not_data", f"the {season} {name} feed has {why}", status=200,
                                         endpoint=FEED_ENDPOINT[name[:3]])
            _save_json(p, rows)
            out[name] = rows
        print(f"  {name}: {len(out[name])} rows{' (cached)' if saved is not None else ''}")
    return out


def stop_refused(e, league):
    """Print why StatsPlus refused, then exit EXIT_REFUSED."""
    print(f"  {e.user_message(league)}")
    print("  No input CSV was written.")
    sys.exit(EXIT_REFUSED)


def load_pull(path):
    rows = json.load(open(path, encoding="utf-8"))
    if isinstance(rows, dict):
        rows = rows.get("rows") or rows.get("players") or []
    rows = S.translate_rows(rows)
    return {str(r.get("ID")): r for r in rows if r.get("ID")}


# ---------------------------------------------------------------- tab builders

def build_stats(feeds, teams, pull):
    B = aggregate(feeds["bat1"], lambda r: r["player_id"], BAT_SUM)
    B2 = aggregate(feeds["bat2"], lambda r: r["player_id"], ["pa"])
    B3 = aggregate(feeds["bat3"], lambda r: r["player_id"], ["pa"])
    P = aggregate(feeds["pit1"], lambda r: r["player_id"], PIT_SUM)
    P2 = aggregate(feeds["pit2"], lambda r: r["player_id"], ["bf"])
    P3 = aggregate(feeds["pit3"], lambda r: r["player_id"], ["bf"])
    F = aggregate(feeds["fld1"], lambda r: (r["player_id"], int(num(r["position"]))), FLD_SUM)

    def org(pid, d):
        pr = pull.get(pid)
        tid = str(pr.get("Org") or pr.get("Team")) if pr else str(d.get("_team"))
        return teams.get(tid, tid)

    def name(pid, d):
        pr = pull.get(pid)
        return pr.get("Name") if pr else f"ID {pid}"

    hit = []
    for pid, d in B.items():
        if d["pa"] <= 0:
            continue
        hit.append([int(pid), name(pid, d), org(pid, d), d["g"], d["gs"], d["pa"], d["ab"], d["h"],
                    d["h"] - d["d"] - d["t"] - d["hr"], d["d"], d["t"], d["hr"], d["rbi"], d["r"], d["bb"],
                    d["ibb"], d["hp"], d["sh"], d["sf"], d["ci"], d["k"], d["gdp"],
                    (d["pitches_seen"] / d["pa"]) if d["pa"] else 0.0, d["sb"], d["cs"], d["ubr"]])
    hit.sort(key=lambda r: -r[5])

    pit = []
    for pid, d in P.items():
        if d["bf"] <= 0 and d["outs"] <= 0:
            continue
        pit.append([int(pid), name(pid, d), org(pid, d), d["g"], d["gs"], d["w"], d["l"], thirds(d["outs"]),
                    d["bf"], d["ab"], d["ha"], d["ha"] - d["da"] - d["ta"] - d["hra"], d["da"], d["ta"],
                    d["hra"], d["r"], d["er"], d["bb"], d["iw"], d["k"], d["hp"], d["sh"], d["sf"], d["wp"],
                    d["bk"], d["ci"], d["dp"], d["ir"], d["irs"], d["pi"], d["gb"], d["fb"], d["sb"], d["cs"]])
    pit.sort(key=lambda r: -r[8])

    tables = []
    for code in sorted(FIELD_POS):
        rows = []
        for (pid, pos), d in F.items():
            if pos != code or (d["g"] <= 0 and d["ip"] <= 0 and d["ipf"] <= 0):
                continue
            rows.append([int(pid), name(pid, d), org(pid, d), code, d["g"], d["gs"], d["tc"], d["a"], d["po"],
                         d["e"], d["dp"], d["tp"], d["zr"], d["sba"], d["rto"], f"{int(d['ip'])}.{int(d['ipf'])}",
                         d["pb"], d["er"], d["opps_0"], d["opps_made_0"], d["opps_1"], d["opps_made_1"],
                         d["opps_2"], d["opps_made_2"], d["opps_3"], d["opps_made_3"], d["opps_4"],
                         d["opps_made_4"], d["opps_5"], d["framing"], d["arm"]])
        rows.sort(key=lambda r: -(num(r[15].split(".")[0]) * 3 + num(r[15].split(".")[1])))
        if not rows:
            sys.exit(f"no fielding rows for position {FIELD_POS[code]}")
        tables.append(rows)
    return {"B": B, "B2": B2, "B3": B3, "P": P, "P2": P2, "P3": P3, "F": F,
            "hit": hit, "pit": pit, "fld": tables}


def build_batter_ratings(st, pull):
    vr, vl, missing = [], [], 0
    for row in st["hit"]:
        pid = str(row[0])
        pr = pull.get(pid)
        if not pr:
            missing += 1
            continue
        base = [int(pid), pr.get("POS"), pr.get("Name")]
        rat = [pr.get("B")] + [pr.get(k) for k in BAT_RAT_HDR[5:]]
        pa_r = st["B3"].get(pid, {}).get("pa", 0.0)
        pa_l = st["B2"].get(pid, {}).get("pa", 0.0)
        if pa_r > 0:
            vr.append(base + [pa_r] + rat)
        if pa_l > 0:
            vl.append(base + [pa_l] + rat)
    vr.sort(key=lambda r: -r[3])
    vl.sort(key=lambda r: -r[3])
    return vr, vl, missing


def build_fielding_ratings(st, pull):
    ip_total = {}
    for (pid, pos), d in st["F"].items():
        if pos in FIELD_POS:
            ip_total[pid] = ip_total.get(pid, 0) + int(d["ip"]) * 3 + int(d["ipf"])
    ids = {str(r[0]) for r in st["hit"]} | set(ip_total)
    rows, skipped = [], 0
    for pid in ids:
        pr = pull.get(pid)
        if not pr:
            skipped += 1
            continue
        if str(pr.get("POS")) in PITCHER_POS:
            continue
        rows.append([int(pid), POS_CODE.get(str(pr.get("POS")), 0), pr.get("Name"), thirds(ip_total.get(pid, 0)),
                     feet_inches(pr.get("HT"))] + [pr.get(k) for k in FLD_RAT_HDR[5:]])
    rows.sort(key=lambda r: -num(r[3].split(".")[0]))
    return rows, skipped


def build_pitcher_ratings(st, pull, ids):
    vr, vl, missing = [], [], []
    for pid in ids:
        pr = pull.get(pid)
        if not pr:
            missing.append(pid)
            continue
        base = [int(pid), pr.get("POS"), pr.get("Name")]
        rat = [pr.get("T")] + [pr.get(k) for k in PIT_RAT_HDR[5:]]
        bf_r = st["P3"].get(pid, {}).get("bf", 0.0)
        bf_l = st["P2"].get(pid, {}).get("bf", 0.0)
        if bf_r > 0:
            vr.append(base + [bf_r] + rat)
        if bf_l > 0:
            vl.append(base + [bf_l] + rat)
    vr.sort(key=lambda r: -r[3])
    vl.sort(key=lambda r: -r[3])
    return vr, vl, missing


def reconcile_roles(st, sp_rows, sp_hdr, rp_rows, rp_hdr):
    """Your SP/RP paste vs the API season totals, per pitcher. Only st["P"]
    is read, so another season's totals work the same way ({"P": totals}).
    OK    = the paste equals the season totals
    OVER  = paste exceeds the season total (playoff games, or another season)
    UNKNOWN = an ID with no MLB pitching this season
    UNDER = one role pasted, or fewer innings than the season
    ABSENT = MLB innings, in neither tab
    paste_season() and paste_decision() say what the roles stage does with them."""
    def ip_outs(v):
        v = num(v)
        return int(v) * 3 + int(round((v - int(v)) * 10))
    keys = [("BF", "bf"), ("HA", "ha"), ("HR", "hra"), ("BB", "bb"), ("K", "k"), ("R", "r")]
    pasted = {}
    for hdr, rows, tag in ((sp_hdr, sp_rows, "SP"), (rp_hdr, rp_rows, "RP")):
        col = {h: i for i, h in enumerate(hdr)}
        for r in rows:
            pid = str(int(num(r[col["ID"]])))
            d = pasted.setdefault(pid, {"tags": set(), "outs": 0, **{k: 0.0 for k, _ in keys}})
            d["tags"].add(tag)
            d["outs"] += ip_outs(r[col["IP"]]) if "IP" in col else 0
            for k, _ in keys:
                d[k] += num(r[col[k]]) if k in col else 0.0
    over, unknown, under, ok = [], [], [], 0
    for pid, d in pasted.items():
        api = st["P"].get(pid)
        if not api or (api["bf"] <= 0 and api["outs"] <= 0):
            unknown.append(pid)
            continue
        diffs = [(k, d[k], api[a]) for k, a in keys if abs(d[k] - api[a]) > 1e-9]
        if abs(d["outs"] - api["outs"]) > 0:
            diffs.append(("IP outs", d["outs"], api["outs"]))
        if any(p > t + 1e-9 for _, p, t in diffs):
            over.append((pid, diffs))
        elif diffs:
            under.append((pid, diffs, d["tags"]))
        else:
            ok += 1
    absent = [pid for pid, api in st["P"].items() if (api["bf"] > 0 or api["outs"] > 0) and pid not in pasted]
    return {"ok": ok, "over": over, "unknown": unknown, "under": under, "absent": absent, "pasted": pasted}


# Which season the paste is from. A paste of this season matches most pitchers'
# totals exactly (OOTP's role split may add playoff games to some). A paste of
# another season matches almost none, and many of its IDs did not pitch in MLB
# this season. A wrong season always stops: --accept-paste never skips it.
SEASON_MATCH = 0.5          # share of the pasted pitchers that match a season's totals exactly
OTHER_SEASON_UNKNOWN = 0.1  # share of the pasted IDs with no MLB pitching in the season


def exact_share(rec):
    """Share of the pasted pitchers whose paste equals the season totals."""
    n = len(rec["pasted"])
    return rec["ok"] / n if n else 0.0


def other_seasons(season, in_game_date):
    """The seasons a wrong paste most likely comes from, nearest first: the two
    before `season`, and the two after it when they are already played (only
    with --year). A season is played once the in-game date reaches October."""
    y, m = int(in_game_date[:4]), int(in_game_date[5:7])
    last = y if m >= 10 else y - 1
    return [s for s in (season - 1, season + 1, season - 2, season + 2) if s <= last]


def paste_season(season, rec, others):
    """The season the SP/RP paste is from, or None.
    rec = reconcile_roles() against `season`; others = {year: reconcile_roles()
    against that year} for the other seasons checked (empty when none).
      season   at least SEASON_MATCH of the pasted pitchers match it, or nothing
               points to another season (the other checks list the odd rows)
      a year   another season matches at least SEASON_MATCH, and better
      None     under SEASON_MATCH match this season and at least
               OTHER_SEASON_UNKNOWN of the pasted IDs did not pitch in it, but
               no season checked matches the paste"""
    here = exact_share(rec)
    if here >= SEASON_MATCH:
        return season
    better = [(exact_share(r), -abs(y - season), y) for y, r in others.items()]
    better = [b for b in better if b[0] >= SEASON_MATCH and b[0] > here]
    if better:
        return max(better)[2]
    n = len(rec["pasted"])
    return None if n and len(rec["unknown"]) / n >= OTHER_SEASON_UNKNOWN else season


def paste_decision(season, found, rec, accept_paste):
    """What the roles stage does after the paste check, found = paste_season():
      "wrong season"  stop; --accept-paste never skips it
      "odd"           stop unless --accept-paste: pasted IDs with no MLB
                      pitching this season, or pitchers above their totals all
                      over the paste
      "confirm"       stop unless --accept-paste: pitchers above their totals
                      (OOTP's role split counts playoff games) or MLB pitchers
                      in neither tab
      None            go on"""
    if found != season:
        return "wrong season"
    if accept_paste:
        return None
    if rec["unknown"] or len(rec["over"]) > max(5, 0.05 * len(rec["pasted"])):
        return "odd"
    if rec["over"] or rec["absent"]:
        return "confirm"
    return None


def wrong_season_text(season, found, rec, others):
    """The stop message for a paste that is not this season's."""
    n = len(rec["pasted"])
    if found is not None:
        lines = [f"  STOP: the SP/RP paste is from season {found}, not {season}.",
                 f"  {others[found]['ok']} of {n} pasted pitchers match the {found} totals; "
                 f"{rec['ok']} match the {season} totals."]
    else:
        checked = ", ".join(str(y) for y in sorted(others))
        lines = [f"  STOP: the SP/RP paste is not from season {season}.",
                 f"  Only {rec['ok']} of {n} pasted pitchers match the {season} totals, and "
                 f"{len(rec['unknown'])} did not pitch in MLB in {season}."
                 + (f" It does not match {checked} either." if checked else "")]
    lines.append(f"  Paste the {season} stats into 'SP Data' (as starter) and 'RP Data' (as reliever), "
                 "save, close Excel, then run this again.")
    return "\n".join(lines)


def season_pitching(base, slug, year, refresh=False):
    """{pid: totals} of one season's MLB pitching (the overall feed), from the
    saved feed .cache/metadata_<slug>_<year>/pit1.json, else from StatsPlus
    (saved there the way fetch_season saves it). None when it can not be read:
    this only names the season of a wrong paste, so it never stops the run."""
    cache_dir = os.path.join(HERE, ".cache", f"metadata_{slug}_{year}")
    p = os.path.join(cache_dir, "pit1.json")
    rows = None
    if os.path.exists(p) and not refresh:
        try:
            rows = json.load(open(p, encoding="utf-8"))
        except ValueError:
            rows = None
        if feed_problem("pit1", rows):
            rows = None
    if rows is None:
        try:
            rows = S.fetch_pitching(base, year=year, split=1)
        except Exception as e:          # noqa: BLE001  refused, network: that season is not checked
            print(f"  (the {year} pitching could not be read: {type(e).__name__}; {year} is not checked)")
            return None
        if feed_problem("pit1", rows):
            return None
        os.makedirs(cache_dir, exist_ok=True)
        _save_json(p, rows)
    return aggregate(rows, lambda r: r["player_id"], PIT_SUM)


def find_paste_season(base, slug, season, date, rec, sp, rp, refresh=False):
    """(paste_season(), {year: reconcile result}). The other seasons are read
    only when the paste does not match `season`. sp / rp = (rows, header)."""
    if exact_share(rec) >= SEASON_MATCH:
        return season, {}
    others = {}
    for y in other_seasons(season, date):
        totals = season_pitching(base, slug, y, refresh)
        if totals is not None:
            others[y] = reconcile_roles({"P": totals}, sp[0], sp[1], rp[0], rp[1])
    return paste_season(season, rec, others), others


# ---------------------------------------------------------------- driver

SEASON_END = "10-01"    # month-day: the regular season is over by this in-game day


def season_end_pull(league, season):
    """The raw file of the league's last live pull on or before the end of the
    season's regular season (in-game {season}-10-01), from the ratings archive.
    The season's stats are paired with the ratings the players had when they
    played, not with an offseason pull (user, 2026-10-01: "end of regular season
    ratings"). None when the archive has no such pull or its file is gone; the
    caller then falls back to the newest pull."""
    import sqlite3
    db = os.path.join(REPO, "tgs-viz", "backtest", "ratings_history.db")
    if not os.path.exists(db):
        return None
    conn = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    try:
        row = conn.execute(
            "SELECT game_date, source_files FROM pulls WHERE league=? AND source='live' "
            "AND game_date BETWEEN ? AND ? ORDER BY game_date DESC, pull_id DESC LIMIT 1",
            (league, f"{season}-01-01", f"{season}-{SEASON_END}")).fetchone()
    finally:
        conn.close()
    if not row:
        print(f"  NOTE: no {league} live pull on or before {season}-{SEASON_END} in the archive; "
              f"the newest pull's ratings are used")
        return None
    try:
        files = json.loads(row[1] or "[]")
    except ValueError:
        files = []
    for rel in files:
        path = rel if os.path.isabs(rel) else os.path.join(REPO, rel)
        if path.endswith(".json") and os.path.exists(path):
            print(f"  ratings: the last pull of the regular season (game date {row[0]})")
            return path
    print(f"  NOTE: the {league} pull of game date {row[0]} has no raw file left; the newest pull's "
          f"ratings are used")
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--league", required=True, choices=list(LEAGUE_DIRS))
    ap.add_argument("--out", required=True, help="directory for the nine CSVs (+ manifest.json)")
    ap.add_argument("--stage", default="all", choices=["auto", "roles", "all"])
    ap.add_argument("--year", type=int, help="season to build (default: from the in-game date)")
    ap.add_argument("--slug", help="StatsPlus league slug (default per league)")
    ap.add_argument("--pull", help="ratings pull cache JSON (default .cache/statsplus_<slug>.json)")
    ap.add_argument("--workbook", help="25 Metadata.xlsx to read pasted tabs from (default the league's)")
    ap.add_argument("--from-workbook", default="SP Data,RP Data",
                    help="comma list of tabs copied from the workbook instead of built "
                         "(SP Data / RP Data always; add ratings tabs only for validation)")
    ap.add_argument("--allow-partial", action="store_true", help="build mid-season anyway")
    ap.add_argument("--accept-paste", "--allow-absent", dest="accept_paste", action="store_true",
                    help="use the SP/RP paste as it is: a few pitchers above their regular-season "
                         "totals (playoff games), pitchers left out of both tabs and IDs with no MLB "
                         "pitching. A paste of another season still stops.")
    ap.add_argument("--refresh", action="store_true", help="refetch the API feeds (ignore cache)")
    a = ap.parse_args()

    if a.year:
        check_engine_boundary(a.league, a.year)     # before any request
    slug = a.slug or SLUGS.get(a.league) or ST.slug(a.league)
    base = S.normalize_base(slug)
    try:
        date = S.fetch_date(base)
    except S.StatsPlusRefused as e:
        stop_refused(e, a.league)
    season = season_of(date, a.year, a.allow_partial)
    check_engine_boundary(a.league, season)         # after /date, before /teams and the stat feeds
    pull_path = a.pull or season_end_pull(a.league, season) or os.path.join(HERE, ".cache", f"statsplus_{slug}.json")
    xlsx = a.workbook or os.path.join(REPO, LEAGUE_DIRS[a.league], "25 Metadata.xlsx")
    from_wb = {t.strip() for t in a.from_workbook.split(",") if t.strip()} | set(PASTE_TABS)
    os.makedirs(a.out, exist_ok=True)
    mpath = os.path.join(a.out, "manifest.json")
    manifest = json.load(open(mpath, encoding="utf-8")) if os.path.exists(mpath) else {}
    if manifest and manifest.get("season") != season and a.stage == "roles":
        sys.exit(f"{a.out} was built for season {manifest.get('season')}, not {season}; rerun --stage auto")

    print(f"== {a.league} metadata inputs: season {season} (in-game date {date}), {base}")
    cache_dir = os.path.join(HERE, ".cache", f"metadata_{slug}_{season}")
    try:
        # /teams first, while the /date read above is still in memory. cache=True: the
        # second bat stage reuses this reply while the in-game date is unchanged.
        team_rows = S.fetch_teams(base, cache=True)
        feeds = fetch_season(base, season, cache_dir, a.refresh)
    except S.StatsPlusRefused as e:
        stop_refused(e, a.league)
    teams = {}
    for t in team_rows:
        teams[str(t.get("ID"))] = f"{t.get('Name', '')} {t.get('Nickname', '')}".strip()
    if not os.path.exists(pull_path):
        sys.exit(f"ratings pull cache missing: {pull_path} (run Get StatsPlus Ratings first)")
    pull = load_pull(pull_path)
    print(f"  ratings pull: {len(pull)} players from {os.path.basename(pull_path)} "
          f"(file date {os.path.getmtime(pull_path):.0f})")
    st = build_stats(feeds, teams, pull)
    lids = sorted({r.get("league_id") for r in feeds["pit1"]})
    manifest.update({"league": a.league, "slug": slug, "season": season, "in_game_date": date,
                     "league_ids": lids, "pull": pull_path, "counts": manifest.get("counts", {})})

    if a.stage in ("auto", "all"):
        write_csv(os.path.join(a.out, "Hitting_Data.csv"), HIT_HDR, st["hit"])
        write_csv(os.path.join(a.out, "Pitching_Data.csv"), PIT_HDR, st["pit"])
        write_fielding(os.path.join(a.out, "Fielding_Data.csv"), FLD_HDR, st["fld"])
        manifest["counts"].update({"Hitting_Data": len(st["hit"]), "Pitching_Data": len(st["pit"]),
                                   "Fielding_Data": [len(t) for t in st["fld"]]})
        print(f"  Hitting_Data {len(st['hit'])} | Pitching_Data {len(st['pit'])} | "
              f"Fielding_Data {[len(t) for t in st['fld']]} (C..RF)")
        if "Batter Ratings" in from_wb:
            write_workbook_tab(xlsx, "Batter Ratings", os.path.join(a.out, "Batter_Ratings.csv"))
            print("  Batter_Ratings: copied from the workbook")
        else:
            vr, vl, miss = build_batter_ratings(st, pull)
            write_side_by_side(os.path.join(a.out, "Batter_Ratings.csv"), BAT_RAT_HDR, vr, vl)
            manifest["counts"].update({"Batter_Ratings_vR": len(vr), "Batter_Ratings_vL": len(vl),
                                       "hitters_not_in_pull": miss})
            print(f"  Batter_Ratings vR {len(vr)} / vL {len(vl)}; hitters absent from the pull "
                  f"(no ratings row, stats still counted): {miss}")
        if "Fielding Ratings" in from_wb:
            write_workbook_tab(xlsx, "Fielding Ratings", os.path.join(a.out, "Fielding_Ratings.csv"))
            print("  Fielding_Ratings: copied from the workbook")
        else:
            fr, skipped = build_fielding_ratings(st, pull)
            write_csv(os.path.join(a.out, "Fielding_Ratings.csv"), FLD_RAT_HDR, fr)
            manifest["counts"].update({"Fielding_Ratings": len(fr), "fielders_not_in_pull": skipped})
            print(f"  Fielding_Ratings {len(fr)} position players; absent from the pull: {skipped}")
        json.dump(manifest, open(mpath, "w", encoding="utf-8"), indent=1)
        if a.stage == "auto":
            print("\n  Next: paste the pitchers' AS-STARTER stats into the 'SP Data' tab and the\n"
                  "  AS-RELIEVER stats into the 'RP Data' tab of\n"
                  f"    {xlsx}\n"
                  "  (StatsPlus has no stats split by role). Save and close Excel, then continue.")
            return

    # ---- roles stage: paste tabs + pitcher ratings lists
    sp_hdr, sp_rows = read_workbook_tab(xlsx, "SP Data")
    rp_hdr, rp_rows = read_workbook_tab(xlsx, "RP Data")
    if not sp_rows or not rp_rows:
        sys.exit("'SP Data' / 'RP Data' in the workbook are empty - paste them first")
    rec = reconcile_roles(st, sp_rows, sp_hdr, rp_rows, rp_hdr)
    print(f"  paste check vs API {season}: {rec['ok']} pitchers match exactly, "
          f"{len(rec['under'])} partial, {len(rec['absent'])} with MLB innings in neither tab")
    if rec["unknown"]:
        print(f"  !! {len(rec['unknown'])} pasted IDs have no MLB pitching in {season}: "
              f"{rec['unknown'][:8]}{' ...' if len(rec['unknown']) > 8 else ''}")
    for pid, diffs in rec["over"][:8]:
        print(f"  !! {pid} is above his regular-season totals: "
              + ", ".join(f"{k} {p:g} vs {t:g}" for k, p, t in diffs if p > t))
    # A paste of another season always stops (the Recalibrate task passes
    # --accept-paste). A few pitchers above their totals (OOTP's role split
    # counts playoff games), pitchers left out of both tabs and a few IDs with
    # no MLB pitching are the user's call: --accept-paste takes them as they are.
    found, others = find_paste_season(base, slug, season, date, rec, (sp_rows, sp_hdr), (rp_rows, rp_hdr),
                                      a.refresh)
    decision = paste_decision(season, found, rec, a.accept_paste)
    if decision == "wrong season":
        sys.exit(wrong_season_text(season, found, rec, others))
    if decision == "odd":
        sys.exit(f"  STOP: parts of the SP/RP paste do not match the {season} totals (see above). "
                 "Rerun with --accept-paste to use it as it is.")
    for pid, diffs, tags in rec["under"][:6]:
        print(f"     partial {pid} ({'+'.join(sorted(tags))}): "
              + ", ".join(f"{k} {p:g} of {t:g}" for k, p, t in diffs[:3]))
    if rec["absent"]:
        top = sorted(rec["absent"], key=lambda p: -st["P"][p]["bf"])[:6]
        share = sum(st["P"][p]["bf"] for p in rec["absent"]) / max(1.0, sum(d["bf"] for d in st["P"].values()))
        print(f"     in neither tab: {len(rec['absent'])} pitchers, {share:.0%} of league BF; largest: "
              + ", ".join(f"{p} ({st['P'][p]['bf']:g} BF)" for p in top))
    if decision == "confirm":
        sys.exit(f"  STOP for your OK: {len(rec['over'])} pitchers above their regular-season totals "
                 f"(playoff games in OOTP's role split), {len(rec['absent'])} with MLB innings in neither "
                 f"tab. Fix the paste, or rerun with --accept-paste to use it as it is.")
    write_workbook_tab(xlsx, "SP Data", os.path.join(a.out, "SP_Data.csv"))
    write_workbook_tab(xlsx, "RP Data", os.path.join(a.out, "RP_Data.csv"))
    sp_ids = [str(int(num(r[sp_hdr.index("ID")]))) for r in sp_rows]
    rp_ids = [str(int(num(r[rp_hdr.index("ID")]))) for r in rp_rows]
    for tab, ids in (("SP Ratings", sp_ids), ("RP Ratings", rp_ids)):
        fn = DERIVED_TABS[tab]
        if tab in from_wb:
            write_workbook_tab(xlsx, tab, os.path.join(a.out, fn))
            print(f"  {fn}: copied from the workbook")
            continue
        vr, vl, miss = build_pitcher_ratings(st, pull, ids)
        write_side_by_side(os.path.join(a.out, fn), PIT_RAT_HDR, vr, vl)
        manifest["counts"].update({f"{tab.replace(' ', '_')}_vR": len(vr), f"{tab.replace(' ', '_')}_vL": len(vl)})
        print(f"  {fn}: vR {len(vr)} / vL {len(vl)} from the pull for your {tab.split()[0]} list "
              f"({len(ids)} pitchers; not in the pull: {len(miss)})")
    manifest["counts"].update({"SP_Data": len(sp_rows), "RP_Data": len(rp_rows)})
    manifest["paste_check"] = {"exact": rec["ok"], "partial": len(rec["under"]), "absent": len(rec["absent"])}
    json.dump(manifest, open(mpath, "w", encoding="utf-8"), indent=1)
    print(f"  wrote {a.out} (manifest.json has the counts)")


if __name__ == "__main__":
    main()
